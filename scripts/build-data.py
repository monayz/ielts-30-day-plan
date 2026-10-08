#!/usr/bin/env python3
"""
Converts data/IELTS_30_Day_Plan.xlsx into data/plan.json and data/prompts.json
for the website to fetch at runtime.

The spreadsheet is the source of truth. This script does not invent or
reword any content -- it only reshapes rows into the JSON the site expects.

Re-run after editing the spreadsheet:
    python3 scripts/build-data.py
"""
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip3 install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "data" / "IELTS_30_Day_Plan.xlsx"
PLAN_OUT = ROOT / "data" / "plan.json"
PROMPTS_OUT = ROOT / "data" / "prompts.json"

# Item Type (spreadsheet) -> step type used by the site's TYPES metadata
ITEM_TYPE_MAP = {
    "overview_video": "overview",
    "video": "lesson",
    "practice": "practice",
    "sample": "sample",
    "ai_prompt": "ai",
    "review": "review",
}


def cell(row, header_index, name):
    return row[header_index[name]]


def clean_url(raw):
    if raw is None:
        return None
    s = str(raw).strip()
    if not s or s == "—":  # em dash means "no link"
        return None
    return s


def clean_text(raw):
    if raw is None:
        return None
    s = str(raw).strip()
    return s or None


def main():
    if not XLSX_PATH.exists():
        sys.exit(f"Spreadsheet not found: {XLSX_PATH}")

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

    # ---- Plan sheet ----
    ws = wb["Plan"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h else "" for h in rows[0]]
    idx = {name: i for i, name in enumerate(header)}

    required = ["Day", "Week", "Skill", "Item #", "Item Type", "Title", "URL", "Prompt ID", "Instructions for Student", "Timer (min)"]
    missing = [c for c in required if c not in idx]
    if missing:
        sys.exit(f"Plan sheet is missing expected column(s): {missing}")

    plan = {}
    unmapped_types = set()
    for row in rows[1:]:
        if row[idx["Day"]] is None:
            continue
        day = int(cell(row, idx, "Day"))
        item_num = cell(row, idx, "Item #")
        raw_type = clean_text(cell(row, idx, "Item Type"))
        mapped_type = ITEM_TYPE_MAP.get(raw_type)
        if mapped_type is None:
            unmapped_types.add(raw_type)
            mapped_type = raw_type  # keep going; surfaced as a warning below

        timer = cell(row, idx, "Timer (min)")
        timer = int(timer) if isinstance(timer, (int, float)) else None

        entry = {
            "skill": clean_text(cell(row, idx, "Skill")),
            "type": mapped_type,
            "title": clean_text(cell(row, idx, "Title")),
            "url": clean_url(cell(row, idx, "URL")),
            "videoUrl": clean_url(cell(row, idx, "Video URL")) if "Video URL" in idx else None,
            "promptId": clean_text(cell(row, idx, "Prompt ID")),
            "instructions": clean_text(cell(row, idx, "Instructions for Student")),
            "timer": timer,
            "itemNum": int(item_num) if isinstance(item_num, (int, float)) else None,
        }
        plan.setdefault(str(day), []).append(entry)

    if unmapped_types:
        print(f"WARNING: unrecognized Item Type value(s), passed through as-is: {sorted(unmapped_types)}", file=sys.stderr)

    # Keep each day's rows in spreadsheet order (Item # order)
    for day_key in plan:
        plan[day_key].sort(key=lambda r: (r["itemNum"] is None, r["itemNum"]))

    # ---- AI Prompts sheet ----
    ws2 = wb["AI Prompts"]
    rows2 = list(ws2.iter_rows(values_only=True))
    header2 = [str(h).strip() if h else "" for h in rows2[0]]
    idx2 = {name: i for i, name in enumerate(header2)}
    if "Prompt ID" not in idx2 or "Prompt Text (copy-ready)" not in idx2:
        sys.exit(f"AI Prompts sheet is missing expected columns. Found: {header2}")

    prompts = {}
    for row in rows2[1:]:
        pid = clean_text(cell(row, idx2, "Prompt ID"))
        if not pid:
            continue
        prompts[pid] = {
            "text": cell(row, idx2, "Prompt Text (copy-ready)") or "",
            "image": clean_url(cell(row, idx2, "Image")) if "Image" in idx2 else None,
        }

    # ---- Sanity check: every promptId referenced in the plan exists ----
    used_ids = {e["promptId"] for rows_ in plan.values() for e in rows_ if e["promptId"]}
    missing_ids = used_ids - set(prompts.keys())
    if missing_ids:
        print(f"WARNING: Plan references Prompt ID(s) with no entry in 'AI Prompts': {sorted(missing_ids)}", file=sys.stderr)

    days_present = sorted(int(d) for d in plan.keys())
    if days_present != list(range(1, 31)):
        print(f"WARNING: expected Days 1-30, found: {days_present}", file=sys.stderr)

    PLAN_OUT.write_text(json.dumps(plan, ensure_ascii=False, indent=2, sort_keys=False), encoding="utf-8")
    PROMPTS_OUT.write_text(json.dumps(prompts, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")

    total_steps = sum(len(v) for v in plan.values())
    print(f"Wrote {PLAN_OUT.relative_to(ROOT)} ({len(plan)} days, {total_steps} steps)")
    print(f"Wrote {PROMPTS_OUT.relative_to(ROOT)} ({len(prompts)} prompts)")


if __name__ == "__main__":
    main()
